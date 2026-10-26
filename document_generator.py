import re
import json
from pathlib import Path
from datetime import datetime
from io import BytesIO

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image as PDFImage,
    PageBreak,
)

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.util import Inches as PPTInches, Pt as PPTPt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ======================================================
# 1. STORAGE
# ======================================================

BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "nova_generated"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ======================================================
# 2. SAFE FILE NAMES
# ======================================================

def safe_filename(title):
    """Convert a title into a safe filename."""
    name = re.sub(r"[^a-zA-Z0-9_-]+", "_", title.strip())
    name = name.strip("_")[:70]
    return name or "nova_document"


def get_output_path(title, extension):
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{safe_filename(title)}_{timestamp}.{extension}"
    return OUTPUT_DIR / filename


# ======================================================
# 3. CONTENT HELPERS
# ======================================================

def normalize_content(content):
    """
    Accept a string or a structured dictionary.
    Expected dictionary:
    {
        "title": "30-Day Job Preparation Plan",
        "sections": [
            {
                "heading": "Week 1",
                "content": "Revise Python fundamentals."
            }
        ],
        "table": {
            "headers": ["Day", "Task", "Goal"],
            "rows": [
                [1, "Python basics", "Revise syntax"]
            ]
        },
        "chart": {
            "title": "Weekly Progress",
            "labels": ["Week 1", "Week 2"],
            "values": [20, 50],
            "ylabel": "Completion (%)"
        }
    }
    """
    if isinstance(content, dict):
        data = content.copy()
    else:
        data = {
            "title": "NOVA Preparation Document",
            "sections": [
                {"heading": "Preparation Notes", "content": str(content)}
            ],
        }

    data.setdefault("title", "NOVA Preparation Document")
    data.setdefault("sections", [])
    return data


def create_chart(chart_data):
    """Create a PNG chart and return its path."""
    labels = chart_data.get("labels", [])
    values = chart_data.get("values", [])

    if not labels or not values or len(labels) != len(values):
        raise ValueError(
            "Chart requires equally sized, non-empty labels and values."
        )

    try:
        values = [float(value) for value in values]
    except (TypeError, ValueError):
        raise ValueError("Chart values must be numeric.")

    path = get_output_path(
        chart_data.get("title", "NOVA_Chart"), "png"
    )

    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.bar(labels, values)
    ax.set_title(chart_data.get("title", "NOVA Progress Chart"))
    ax.set_ylabel(chart_data.get("ylabel", "Value"))
    ax.tick_params(axis="x", rotation=30)
    fig.tight_layout()
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)

    return path


# ======================================================
# 4. PDF GENERATOR
# ======================================================

def create_pdf(content):
    data = normalize_content(content)
    path = get_output_path(data["title"], "pdf")
    story = []

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="NOVA_Title",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#4338CA"),
        spaceAfter=18,
    ))
    styles.add(ParagraphStyle(
        name="NOVA_Heading",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#4338CA"),
        spaceBefore=12,
        spaceAfter=6,
    ))

    story.append(Paragraph(data["title"], styles["NOVA_Title"]))
    story.append(Paragraph(
        f"Created by NOVA | {datetime.now():%d %B %Y}",
        styles["Normal"],
    ))
    story.append(Spacer(1, 18))

    for section in data.get("sections", []):
        heading = section.get("heading", "")
        text = section.get("content", "")

        if heading:
            story.append(Paragraph(str(heading), styles["NOVA_Heading"]))

        # Preserve line breaks and escape HTML-sensitive characters.
        safe_text = (
            str(text)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
            .replace("\n", "<br/>")
        )
        if safe_text:
            story.append(Paragraph(safe_text, styles["BodyText"]))
            story.append(Spacer(1, 8))

    table_data = data.get("table")
    if table_data:
        headers = table_data.get("headers", [])
        rows = table_data.get("rows", [])

        if headers:
            table = Table(
                [[str(value) for value in headers]]
                + [[str(value) for value in row] for row in rows],
                repeatRows=1,
                hAlign="LEFT",
            )
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4338CA")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1),
                 [colors.white, colors.HexColor("#F3F4F6")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(Spacer(1, 12))
            story.append(table)
            story.append(Spacer(1, 16))

    chart_data = data.get("chart")
    if chart_data:
        chart_path = create_chart(chart_data)
        story.append(Paragraph(
            chart_data.get("title", "Progress Chart"),
            styles["NOVA_Heading"],
        ))
        story.append(PDFImage(str(chart_path), width=6.3 * inch, height=3.3 * inch))

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=0.65 * inch,
        leftMargin=0.65 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=data["title"],
        author="NOVA AI",
    )
    doc.build(story)

    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("PDF creation failed.")

    return path


# ======================================================
# 5. WORD DOCUMENT GENERATOR
# ======================================================

def create_docx(content):
    data = normalize_content(content)
    path = get_output_path(data["title"], "docx")

    doc = Document()
    doc.add_heading(data["title"], 0)
    doc.add_paragraph(
        f"Created by NOVA | {datetime.now():%d %B %Y}"
    )

    for section in data.get("sections", []):
        if section.get("heading"):
            doc.add_heading(section["heading"], level=1)
        if section.get("content"):
            doc.add_paragraph(str(section["content"]))

    table_data = data.get("table")
    if table_data and table_data.get("headers"):
        headers = table_data["headers"]
        rows = table_data.get("rows", [])
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = "Light Shading Accent 1"

        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = str(header)

        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row[:len(headers)]):
                cells[index].text = str(value)

    if data.get("chart"):
        chart_path = create_chart(data["chart"])
        doc.add_picture(str(chart_path), width=Inches(6))

    doc.save(path)

    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("Word document creation failed.")

    return path


# ======================================================
# 6. POWERPOINT GENERATOR
# ======================================================

def create_pptx(content):
    data = normalize_content(content)
    path = get_output_path(data["title"], "pptx")

    prs = Presentation()

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = data["title"]
    slide.placeholders[1].text = (
        f"NOVA AI Preparation Resource\n"
        f"{datetime.now():%d %B %Y}"
    )

    for section in data.get("sections", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = section.get("heading", "Key Points")

        body = slide.placeholders[1].text_frame
        body.clear()

        text = str(section.get("content", ""))
        lines = [line.strip() for line in text.splitlines() if line.strip()]

        if not lines:
            lines = ["No additional details provided."]

        for index, line in enumerate(lines):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = PPTPt(20)
            paragraph.level = 0

    table_data = data.get("table")
    if table_data and table_data.get("headers"):
        headers = table_data["headers"]
        rows = table_data.get("rows", [])

        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = "Schedule / Data"

        table = slide.shapes.add_table(
            len(rows) + 1,
            len(headers),
            PPTInches(0.4),
            PPTInches(1.5),
            PPTInches(12.5),
            PPTInches(5),
        ).table

        for col, header in enumerate(headers):
            table.cell(0, col).text = str(header)

        for row_index, row in enumerate(rows, start=1):
            for col, value in enumerate(row[:len(headers)]):
                table.cell(row_index, col).text = str(value)

    if data.get("chart"):
        chart_path = create_chart(data["chart"])
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = data["chart"].get(
            "title", "Progress Chart"
        )
        slide.shapes.add_picture(
            str(chart_path),
            PPTInches(1),
            PPTInches(1.4),
            width=PPTInches(11),
        )

    prs.save(path)

    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("PowerPoint creation failed.")

    return path


# ======================================================
# 7. EXCEL WORKBOOK GENERATOR
# ======================================================

def create_xlsx(content):
    data = normalize_content(content)
    path = get_output_path(data["title"], "xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "NOVA Preparation"

    ws["A1"] = data["title"]
    ws["A1"].font = Font(size=18, bold=True, color="4338CA")
    ws["A2"] = f"Created by NOVA | {datetime.now():%d %B %Y}"

    row_number = 4

    for section in data.get("sections", []):
        heading = section.get("heading", "")
        text = section.get("content", "")

        if heading:
            ws.cell(row=row_number, column=1, value=heading)
            ws.cell(row=row_number, column=1).font = Font(bold=True, size=13)
            row_number += 1

        if text:
            ws.cell(row=row_number, column=1, value=str(text))
            ws.cell(row=row_number, column=1).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            row_number += 2

    table_data = data.get("table")
    if table_data and table_data.get("headers"):
        ws = wb.create_sheet("Schedule")
        headers = table_data["headers"]
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid", fgColor="4338CA"
            )
            cell.alignment = Alignment(wrap_text=True)

        for row in table_data.get("rows", []):
            ws.append(list(row))

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for column_cells in ws.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            ws.column_dimensions[
                column_cells[0].column_letter
            ].width = min(max_length + 3, 40)

    wb.save(path)

    if not path.exists() or path.stat().st_size == 0:
        raise RuntimeError("Excel workbook creation failed.")

    return path


# ======================================================
# 8. MAIN GENERATOR
# ======================================================

GENERATORS = {
    "pdf": create_pdf,
    "docx": create_docx,
    "pptx": create_pptx,
    "xlsx": create_xlsx,
}


def generate_document(content, file_format):
    """
    Generate a real file and return its Path.

    Supported formats: pdf, docx, pptx, xlsx.
    """
    file_format = file_format.lower().strip().lstrip(".")

    if file_format not in GENERATORS:
        raise ValueError(
            "Supported formats are PDF, DOCX, PPTX, and XLSX."
        )

    path = GENERATORS[file_format](content)

    # Confirm the file is available before returning it.
    if not path.is_file():
        raise RuntimeError("The generated file could not be found.")

    return path
